"""
Convert Chandra OCR HTML output → DOCX.
Direct HTML element mapping, no zone classification.
Embeds stamp/seal images by cropping from original PDF.
"""
import io
import re
import logging
from typing import Optional

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
from PIL import Image

# Page dimensions for orientation switching
_A4_LONG = Cm(29.7)
_A4_SHORT = Cm(21.0)

logger = logging.getLogger("HtmlToDocx")

# NĐ 30/2020 formatting
FONT_NAME = "Times New Roman"
FONT_SIZE = Pt(14)
MARGIN_TOP = Cm(2)
MARGIN_BOTTOM = Cm(2)
MARGIN_LEFT = Cm(3)
MARGIN_RIGHT = Cm(2)


def _get_alignment(style: str) -> Optional[int]:
    """Extract text-align from inline style."""
    m = re.search(r'text-align:\s*(center|right|left|justify)', style or '')
    if not m:
        return None
    return {
        'center': WD_ALIGN_PARAGRAPH.CENTER,
        'right': WD_ALIGN_PARAGRAPH.RIGHT,
        'left': WD_ALIGN_PARAGRAPH.LEFT,
        'justify': WD_ALIGN_PARAGRAPH.JUSTIFY,
    }.get(m.group(1))


def _add_runs(paragraph, element, parent_bold=False, parent_italic=False, parent_underline=False):
    """Recursively add text runs from HTML element, preserving formatting."""
    for child in element.children:
        if isinstance(child, NavigableString):
            text = str(child)
            if text.strip() or text == ' ':
                run = paragraph.add_run(text)
                run.font.name = FONT_NAME
                run.font.size = FONT_SIZE
                run.bold = parent_bold
                run.italic = parent_italic
                run.underline = parent_underline
        elif isinstance(child, Tag):
            if child.name == 'br':
                run = paragraph.add_run('\n')
                run.font.name = FONT_NAME
                run.font.size = FONT_SIZE
            elif child.name in ('b', 'strong'):
                _add_runs(paragraph, child, parent_bold=True,
                          parent_italic=parent_italic, parent_underline=parent_underline)
            elif child.name in ('i', 'em'):
                _add_runs(paragraph, child, parent_bold=parent_bold,
                          parent_italic=True, parent_underline=parent_underline)
            elif child.name == 'u':
                _add_runs(paragraph, child, parent_bold=parent_bold,
                          parent_italic=parent_italic, parent_underline=True)
            elif child.name == 'p':
                # Nested <p> inside <div> — treat as paragraph content
                align = _get_alignment(child.get('style', ''))
                if align is not None:
                    paragraph.alignment = align
                _add_runs(paragraph, child, parent_bold, parent_italic, parent_underline)
            else:
                _add_runs(paragraph, child, parent_bold, parent_italic, parent_underline)


def _remove_borders(cell):
    """Remove all borders from a table cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    existing = tcPr.find(qn('w:tcBorders'))
    if existing is not None:
        tcPr.remove(existing)
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="none" w:sz="0" w:space="0"/>'
        '<w:left w:val="none" w:sz="0" w:space="0"/>'
        '<w:bottom w:val="none" w:sz="0" w:space="0"/>'
        '<w:right w:val="none" w:sz="0" w:space="0"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)


def _apply_margins(section, orientation="portrait"):
    """Set NĐ30 margins + page size for given orientation."""
    section.top_margin = MARGIN_TOP
    section.bottom_margin = MARGIN_BOTTOM
    section.left_margin = MARGIN_LEFT
    section.right_margin = MARGIN_RIGHT
    if orientation == "landscape":
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = _A4_LONG
        section.page_height = _A4_SHORT
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = _A4_SHORT
        section.page_height = _A4_LONG


# Heuristic threshold: a table with > WIDE_TABLE_COL_THRESHOLD effective cols
# is treated as landscape-worthy (A4 portrait fits ~6-8 columns comfortably).
WIDE_TABLE_COL_THRESHOLD = 8


def _effective_table_cols(table_el: Tag) -> int:
    """Return the max row-sum of colspan in a <table> — the effective width.

    Example: a row `<th colspan="5">A</th><th colspan="5">B</th>` → 10 cols,
    even though only 2 <th> elements exist.
    """
    max_cols = 0
    for tr in table_el.find_all('tr'):
        row_sum = 0
        for cell in tr.find_all(['td', 'th']):
            try:
                row_sum += max(1, int(str(cell.get('colspan', 1)).strip() or '1'))
            except (ValueError, TypeError):
                row_sum += 1
        if row_sum > max_cols:
            max_cols = row_sum
    return max_cols


def _split_html_heuristic(html: str):
    """Fallback split when no <!--page-orient--> markers are present.

    Scan top-level elements. If an element contains a `<table>` whose effective
    col count exceeds WIDE_TABLE_COL_THRESHOLD, emit it as a landscape segment;
    surrounding portrait content stays portrait. This preserves visual layout
    for HTMLs that lost markers during QTextEdit.toPlainText() save-edit, or
    were saved before the orientations pipeline was wired.

    Design rule: a portrait-only HTML (no wide tables) returns a single
    portrait segment — behaves identically to the pre-heuristic code path.
    """
    soup = BeautifulSoup(html, 'html.parser')
    root = soup.body or soup

    # Quick check: any wide table at all? If not, portrait-only path.
    wide_tables = [t for t in root.find_all('table')
                   if _effective_table_cols(t) > WIDE_TABLE_COL_THRESHOLD]
    if not wide_tables:
        return [('portrait', html.strip())]

    segments = []
    portrait_buf = []

    def flush_portrait():
        if portrait_buf:
            content = ''.join(str(c) for c in portrait_buf)
            if content.strip():
                segments.append(('portrait', content))
            portrait_buf.clear()

    for child in root.children:
        if not hasattr(child, 'name') or child.name is None:
            # Text nodes / comments — keep in current portrait buffer
            portrait_buf.append(child)
            continue

        # Does this top-level element carry a wide table?
        is_wide = False
        if child.name == 'table' and _effective_table_cols(child) > WIDE_TABLE_COL_THRESHOLD:
            is_wide = True
        else:
            for inner in child.find_all('table'):
                if _effective_table_cols(inner) > WIDE_TABLE_COL_THRESHOLD:
                    is_wide = True
                    break

        if is_wide:
            flush_portrait()
            segments.append(('landscape', str(child)))
        else:
            portrait_buf.append(child)

    flush_portrait()

    if not segments:
        return [('portrait', html.strip())]
    return segments


def _split_html_by_orientation(html: str):
    """Split HTML into (orientation, segment_html) tuples.

    Two modes:
      1. Marker-driven (preferred): HTML contains `<!--page-orient:X-->` markers
         (emitted by parse_chandra_html when orientations param is provided).
         Split on markers; each marker sets orientation for content that follows.
      2. Heuristic fallback: no markers → inspect tables in the DOM. Wide
         tables (effective cols > WIDE_TABLE_COL_THRESHOLD) trigger landscape
         segments; other content stays portrait. Guards against markers being
         stripped by QTextEdit.toPlainText() during user-save, or OCR runs
         that predate the orientations pipeline.
    """
    pattern = re.compile(r'<!--page-orient:(portrait|landscape)-->', re.IGNORECASE)

    if pattern.search(html):
        segments = []
        last_end = 0
        current_orient = "portrait"
        for m in pattern.finditer(html):
            segment = html[last_end:m.start()].strip()
            if segment:
                segments.append((current_orient, segment))
            current_orient = m.group(1).lower()
            last_end = m.end()
        tail = html[last_end:].strip()
        if tail:
            segments.append((current_orient, tail))
        if not segments:
            segments.append(("portrait", html.strip()))
        return segments

    # No markers → heuristic
    return _split_html_heuristic(html)


def _render_segment(doc, segment_html, original_image=None):
    """Render one page's worth of HTML into current section of doc.

    Handles both nested-in-div structure (standard Chandra output) and
    top-level <table>/<p> (common on landscape pages where Chandra emits
    table at root level after a prose preamble).
    """
    soup = BeautifulSoup(segment_html, 'html.parser')
    root = soup.body or soup

    for child in root.children:
        if not hasattr(child, 'name') or child.name is None:
            continue  # Skip NavigableString / comments / text nodes (e.g. prose preamble)

        if child.name == 'table':
            # Top-level table (landscape layout)
            _add_table(doc, child)
            continue

        if child.name == 'p':
            # Top-level paragraph (after prose preamble, before/after table)
            if not child.get_text().strip() and not child.find('img'):
                continue
            para = doc.add_paragraph()
            align = _get_alignment(child.get('style', ''))
            if align is not None:
                para.alignment = align
            _add_runs(para, child)
            continue

        if child.name == 'ul':
            for li in child.find_all('li'):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(li.get_text().strip())
                run.font.name = FONT_NAME
                run.font.size = Pt(11)
            continue

        if child.name != 'div':
            continue  # Unknown top-level tag — skip

        # Standard Chandra structure: div data-label=...
        div = child
        label = div.get('data-label', '')
        bbox_str = div.get('data-bbox', '')

        table_el = div.find('table')
        if table_el:
            _add_table(doc, table_el)
            continue

        img_el = div.find('img')
        if img_el and original_image and bbox_str:
            _add_image(doc, img_el, bbox_str, original_image)
            continue

        ul_el = div.find('ul')
        if ul_el:
            for li in ul_el.find_all('li'):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(li.get_text().strip())
                run.font.name = FONT_NAME
                run.font.size = Pt(11)
            continue

        for p_el in div.find_all('p', recursive=False):
            if not p_el.get_text().strip() and not p_el.find('img'):
                continue
            para = doc.add_paragraph()
            align = _get_alignment(p_el.get('style', ''))
            if align is not None:
                para.alignment = align
            _add_runs(para, p_el)

        if not div.find('p') and not div.find('table') and not div.find('ul'):
            text = div.get_text().strip()
            if text:
                para = doc.add_paragraph()
                _add_runs(para, div)


def convert_html_to_docx(
    html: str,
    output_path: str,
    original_image: Optional[Image.Image] = None,
) -> str:
    """
    Convert Chandra OCR HTML output to DOCX, with per-page orientation.

    Reads <!--page-orient:portrait|landscape--> markers embedded by
    parse_chandra_html. Each page gets its own section with matching
    orientation so printed/exported DOCX shows wide tables landscape.

    Math formulas in <math>…</math> blocks (LaTeX-flavored) are converted
    to readable Unicode text via math_render.strip_math_tags before
    rendering — so the DOCX shows `(A)/(12) × 6 × 70%` rather than raw
    LaTeX `\\frac{A}{12} \\times 6 \\times 70\\%`.

    Args:
        html: Chandra HTML output (possibly with orientation markers)
        output_path: Path to save .docx file
        original_image: Original PDF page image for stamp cropping

    Returns:
        Path to saved DOCX file
    """
    # Clean think tags before segmenting
    html = re.sub(r'<think>.*?</think>', '', html, flags=re.DOTALL).strip()
    # Expand <math>LaTeX</math> to Unicode math text
    from math_render import strip_math_tags
    html = strip_math_tags(html, plain=False)

    # Split by orientation markers (empty docs / no-marker docs → 1 portrait segment)
    segments = _split_html_by_orientation(html)

    doc = Document()
    # Default style
    style = doc.styles['Normal']
    style.font.name = FONT_NAME
    style.font.size = FONT_SIZE
    style.paragraph_format.space_after = Pt(3)
    style.paragraph_format.line_spacing = 1.5

    # First section uses the default one; subsequent sections added per page
    first_orient = segments[0][0]
    _apply_margins(doc.sections[0], first_orient)
    _render_segment(doc, segments[0][1], original_image)

    for orient, seg_html in segments[1:]:
        # Add a new section with NEW_PAGE break, then set its orientation
        new_section = doc.add_section(WD_SECTION.NEW_PAGE)
        _apply_margins(new_section, orient)
        _render_segment(doc, seg_html, original_image)

    doc.save(output_path)
    logger.info(f"DOCX saved: {output_path}")
    return output_path


def _safe_span(val, default=1):
    """Parse rowspan/colspan attribute robustly. Returns ≥1 int."""
    try:
        v = int(str(val).strip())
        return max(1, v)
    except (ValueError, TypeError, AttributeError):
        return default


def _add_table(doc: Document, table_el: Tag):
    """Convert HTML <table> to DOCX table with rowspan/colspan support.

    Two-pass algorithm:
      1. Simulate cell placement on an abstract grid, honoring rowspan/colspan
         — produces a list of placements and the total (rows, cols) shape.
      2. Create a DOCX table with that shape, place each cell at its logical
         origin, then merge the covered rectangle when rs>1 or cs>1.

    Flat tables (all rs=cs=1) reduce exactly to the old single-pass behavior:
    occupied tracking is a no-op, ci increments by 1 per cell, no merges happen.

    Merged headers (e.g. `<th colspan=5>Lương hiện hưởng</th>`) previously
    rendered as a single cell in column 5 with 4 trailing empty columns —
    visually breaking the header row. This fix places the text once and
    merges cells 5–9 so the label correctly spans its child columns.
    """
    all_trs = table_el.find_all('tr')
    if not all_trs:
        return

    # Identify header rows: either inside <thead>, or consisting entirely of <th>
    header_row_indices = set()
    thead = table_el.find('thead')
    if thead:
        thead_trs = thead.find_all('tr')
        for ri, tr in enumerate(all_trs):
            if tr in thead_trs:
                header_row_indices.add(ri)
    for ri, tr in enumerate(all_trs):
        cells = tr.find_all(['td', 'th'])
        if cells and all(c.name == 'th' for c in cells):
            header_row_indices.add(ri)

    # Pass 1: abstract placement on (row, col) grid
    occupied = set()
    placements = []  # (ri, ci, rs, cs, text, is_header)
    max_col = 0
    for ri, tr in enumerate(all_trs):
        ci = 0
        for cell in tr.find_all(['td', 'th']):
            # Skip positions already spanned by prior rowspan
            while (ri, ci) in occupied:
                ci += 1
            rs = _safe_span(cell.get('rowspan', 1))
            cs = _safe_span(cell.get('colspan', 1))
            text = cell.get_text(strip=True)
            is_header = (ri in header_row_indices) or (cell.name == 'th')
            placements.append((ri, ci, rs, cs, text, is_header))
            for r in range(ri, ri + rs):
                for c in range(ci, ci + cs):
                    occupied.add((r, c))
            ci += cs
        if ci > max_col:
            max_col = ci

    total_rows = len(all_trs)
    total_cols = max_col
    if occupied:
        # Also account for rowspan extending beyond explicit row's width
        total_cols = max(total_cols, max(c for (_, c) in occupied) + 1)
    if total_rows == 0 or total_cols == 0:
        return

    # Pass 2: build DOCX table and place cells
    table = doc.add_table(rows=total_rows, cols=total_cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for (ri, ci, rs, cs, text, is_header) in placements:
        # Clamp span to table bounds (defensive — malformed HTML)
        rs = min(rs, total_rows - ri)
        cs = min(cs, total_cols - ci)
        if rs < 1 or cs < 1:
            continue
        try:
            top_left = table.cell(ri, ci)
        except IndexError:
            continue

        if rs > 1 or cs > 1:
            try:
                bottom_right = table.cell(ri + rs - 1, ci + cs - 1)
                target = top_left.merge(bottom_right)
            except Exception as e:  # pragma: no cover — defensive
                logger.debug(f"merge failed at ({ri},{ci}) rs={rs} cs={cs}: {e}")
                target = top_left
        else:
            target = top_left

        # Set content. cell.text='' clears default paragraph, but .merge() may
        # already have content from adjacent cells — so we iterate paragraphs
        # and keep only one with our text.
        for para in target.paragraphs[1:]:
            p_el = para._element
            p_el.getparent().remove(p_el)
        para = target.paragraphs[0]
        for run in list(para.runs):
            run._element.getparent().remove(run._element)
        run = para.add_run(text)
        run.font.name = FONT_NAME
        run.font.size = Pt(12)
        if is_header:
            run.bold = True
            shading = parse_xml(
                f'<w:shd {nsdecls("w")} w:fill="F0F0F0" w:val="clear"/>'
            )
            target._tc.get_or_add_tcPr().append(shading)


def _add_image(doc: Document, img_el: Tag, bbox_str: str, original_image: Image.Image):
    """Crop stamp/seal from original image using bbox and embed in DOCX."""
    try:
        coords = [int(x) for x in bbox_str.split()]
        if len(coords) != 4:
            return

        x1, y1, x2, y2 = coords
        img_w, img_h = original_image.size

        # Scale bbox from Chandra coords (relative to ~1000px) to actual image
        # Chandra uses coordinates relative to the rendered image
        # The actual image might be different size
        scale_x = img_w / 1000 if x2 < 1100 else 1
        scale_y = img_h / 1000 if y2 < 1100 else 1

        crop_x1 = max(0, int(x1 * scale_x))
        crop_y1 = max(0, int(y1 * scale_y))
        crop_x2 = min(img_w, int(x2 * scale_x))
        crop_y2 = min(img_h, int(y2 * scale_y))

        cropped = original_image.crop((crop_x1, crop_y1, crop_x2, crop_y2))

        # Convert to bytes
        buf = io.BytesIO()
        cropped.save(buf, format='PNG')
        buf.seek(0)

        # Calculate DOCX width (max 5cm for stamp)
        aspect = cropped.width / cropped.height
        width = min(Cm(5), Cm(5 * aspect))

        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run()
        run.add_picture(buf, width=width)

    except Exception as e:
        logger.warning(f"Failed to embed image: {e}")
