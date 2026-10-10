import tempfile
import unittest
from pathlib import Path

from docx import Document

from chandra_field_parser import parse_chandra_html
from display_html_renderer import build_display_html
from html_to_docx import convert_html_to_docx


ADMIN_HTML = """
<div data-label="Text" data-bbox="20 20 430 100"><p>BỘ NỘI VỤ<br/>CỤC VĂN THƯ</p></div>
<div data-label="Text" data-bbox="560 20 980 100"><p>CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM<br/>Độc lập - Tự do - Hạnh phúc</p></div>
<div data-label="Text" data-bbox="20 150 430 200"><p>Số: 123/CV-CVT</p></div>
<div data-label="Text" data-bbox="560 150 980 200"><p>Hà Nội, ngày 11 tháng 7 năm 2026</p></div>
<div data-label="Section-Header" data-bbox="250 280 750 330"><p><b>THÔNG BÁO</b></p></div>
<div data-label="Section-Header" data-bbox="180 340 820 400"><p><b>Về việc kiểm thử hệ thống OCR</b></p></div>
<div data-label="Table" data-bbox="80 450 920 700"><table><tr><td>Ô A</td><td>Ô B</td></tr></table></div>
<div data-label="List-Group" data-bbox="20 750 430 850"><p>Nơi nhận:</p><ul><li>Đơn vị A</li><li>Lưu VT</li></ul></div>
<div data-label="Text" data-bbox="600 740 940 780"><p><b>GIÁM ĐỐC</b></p></div>
<div data-label="Image" data-bbox="650 790 850 900"><img alt="signature"/></div>
<div data-label="Text" data-bbox="620 910 940 950"><p><b>Nguyễn Văn A</b></p></div>
"""


class ParserAndExportTests(unittest.TestCase):
    def test_parser_extracts_core_administrative_fields(self):
        result = parse_chandra_html([ADMIN_HTML], page_widths=[1000])

        self.assertEqual("BỘ NỘI VỤ", result["CQBH_tren"])
        self.assertEqual("CỤC VĂN THƯ", result["CQBH_duoi"])
        self.assertEqual("123/CV-CVT", result["So_Ki_Hieu"])
        self.assertEqual("11/07/2026", result["Ngay_BH"])
        self.assertEqual("Thông báo", result["Loai_VB"])
        self.assertIn("kiểm thử hệ thống OCR", result["ND_Chinh"])

    def test_display_and_docx_export_preserve_table_cells(self):
        rendered = build_display_html(ADMIN_HTML)
        self.assertIn("<table", rendered)
        self.assertIn("Ô A", rendered)

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "document.docx"
            convert_html_to_docx(ADMIN_HTML, str(output))
            document = Document(str(output))

        self.assertGreaterEqual(len(document.tables), 1)
        cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        self.assertIn("Ô A", cells)
        self.assertIn("Ô B", cells)


if __name__ == "__main__":
    unittest.main()
