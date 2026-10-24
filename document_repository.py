"""SQLite repository and persistence services for administrative documents."""

from contextlib import closing
from datetime import datetime
import json
import logging
import os
from pathlib import Path
import shutil
import sqlite3
import sys
import threading
import time
from typing import Any, Dict, List, Optional, Union

import fitz
import pandas as pd

from ai_search import VectorStore
from app_config import ACTIVE_OCR_MODEL


BASE_DIR = Path(__file__).parent
DATABASE_DIR = BASE_DIR / "database"
BACKUP_DIR = BASE_DIR / "backup"
DATABASE_TIMEOUT = 30.0
SCHEMA_VERSION = 1
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tiff", ".bmp"}

_HERE = os.path.dirname(os.path.abspath(__file__))
POPPLER_PATH = ""
for _candidate in (
    os.path.join(_HERE, "tools", "poppler", "Library", "bin"),
    os.path.join(_HERE, "..", "tools", "poppler", "Library", "bin"),
):
    if os.path.isdir(_candidate):
        POPPLER_PATH = _candidate
        break

BACKUP_DIR.mkdir(exist_ok=True, parents=True)
logger = logging.getLogger("OCRApp")



class DBConnectionPool:
    """Thread-safe database connection pool for SQLite"""
    
    def __init__(self, db_path, max_connections=10, timeout=30.0):
        self.db_path = db_path
        self.max_connections = max_connections
        self.timeout = timeout
        self.connections = {}
        self.lock = threading.RLock()
        
    def get_connection(self):
        """Get a connection for the current thread"""
        thread_id = threading.get_ident()
        
        with self.lock:
            if thread_id not in self.connections:
                if len(self.connections) >= self.max_connections:
                    # Find and close the oldest connection
                    oldest_thread = min(self.connections.keys(), 
                                       key=lambda k: self.connections[k]['last_used'])
                    self.connections[oldest_thread]['conn'].close()
                    del self.connections[oldest_thread]
                
                # Create new connection
                conn = sqlite3.connect(self.db_path, timeout=self.timeout)
                conn.row_factory = sqlite3.Row
                self._optimize_connection(conn)
                
                self.connections[thread_id] = {
                    'conn': conn,
                    'last_used': time.time()
                }
            else:
                # Update last used time
                self.connections[thread_id]['last_used'] = time.time()
                
            return self.connections[thread_id]['conn']
    
    def _optimize_connection(self, conn):
        """Optimize SQLite connection settings"""
        conn.execute("PRAGMA journal_mode = WAL")  # Write-Ahead Logging
        conn.execute("PRAGMA synchronous = NORMAL")  # Faster writes, still safe
        conn.execute("PRAGMA busy_timeout = 30000")  # 30 second timeout
        conn.execute("PRAGMA temp_store = MEMORY")  # Store temp data in memory
        conn.execute("PRAGMA foreign_keys = ON")    # Enable foreign key constraints
        return conn
    
    def close_all(self):
        """Close all connections in the pool"""
        with self.lock:
            for conn_data in self.connections.values():
                try:
                    conn_data['conn'].close()
                except:
                    pass
            self.connections.clear()



class DocumentDatabase:
    """Database manager for documents and OCR results"""
    
    def __init__(self, db_path=DATABASE_DIR / "documents.db"):
        self.db_path = db_path
        self.suggestions_cache = {}
        self.conn_pool = DBConnectionPool(db_path, max_connections=10, timeout=DATABASE_TIMEOUT)
        self.init_db()
        self.load_suggestions()

    def init_db(self):
        """Initialize database with tables"""
        with closing(sqlite3.connect(self.db_path)) as conn, conn:
            cursor = conn.cursor()
            
            # Create documents table first with all required columns
            cursor.execute('''
                CREATE TABLE IF NOT EXISTS documents (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    file_path TEXT NOT NULL,
                    file_name TEXT NOT NULL,
                    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                    last_modified TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                    file_hash TEXT,
                    file_size INTEGER,
                    page_count INTEGER
                )
            ''')
            
            # Document versions table - Bao gồm field van_ban_dang_txt cho toàn bộ text OCR
            cursor.execute('''
                CREATE TABLE IF NOT EXISTS document_versions (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    document_id INTEGER,
                    version_number INTEGER,
                    cqbh_tren TEXT,
                    cqbh_duoi TEXT,
                    so_ki_hieu TEXT,
                    ngay_bh TEXT,
                    do_khan TEXT,
                    do_mat TEXT,
                    loai_vb TEXT,
                    nd_chinh TEXT,
                    noi_nhan TEXT,
                    chuc_vu TEXT,
                    chu_ky TEXT,
                    van_ban_dang_txt TEXT,
                    modified_by TEXT,
                    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                    FOREIGN KEY (document_id) REFERENCES documents(id)
                )
            ''')
            
            # Schema v1: persist the actual OCR model for every version.
            version_columns = {
                row[1] for row in cursor.execute("PRAGMA table_info(document_versions)")
            }
            if "model_used" not in version_columns:
                cursor.execute(
                    "ALTER TABLE document_versions "
                    "ADD COLUMN model_used TEXT DEFAULT 'chandra-q4'"
                )
            cursor.execute(
                "UPDATE document_versions SET model_used = ? "
                "WHERE model_used IS NULL OR TRIM(model_used) = ''",
                (ACTIVE_OCR_MODEL,),
            )
            cursor.execute(f"PRAGMA user_version = {SCHEMA_VERSION}")

            # Page detections table
            cursor.execute('''
                CREATE TABLE IF NOT EXISTS page_detections (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    document_id INTEGER,
                    page_number INTEGER,
                    detection_data TEXT,
                    page_text TEXT,
                    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                    FOREIGN KEY (document_id) REFERENCES documents(id)
                )
            ''')
            
            # Field suggestions table
            cursor.execute('''
                CREATE TABLE IF NOT EXISTS field_suggestions (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    field_name TEXT NOT NULL,
                    value TEXT NOT NULL,
                    frequency INTEGER DEFAULT 1,
                    last_used TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                    UNIQUE(field_name, value)
                )
            ''')
            
            # Document backup table
            cursor.execute('''
                CREATE TABLE IF NOT EXISTS document_backups (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    document_id INTEGER,
                    backup_path TEXT NOT NULL,
                    reason TEXT,
                    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                    FOREIGN KEY (document_id) REFERENCES documents(id)
                )
            ''')
            
            # Document tags table
            cursor.execute('''
                CREATE TABLE IF NOT EXISTS document_tags (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,

                    document_id INTEGER,
                    tag_name TEXT NOT NULL,
                    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                    FOREIGN KEY (document_id) REFERENCES documents(id),
                    UNIQUE(document_id, tag_name)
                )
            ''')
            
            # Add indexes for performance
            cursor.execute('CREATE INDEX IF NOT EXISTS idx_doc_filename ON documents(file_name)')
            cursor.execute('CREATE INDEX IF NOT EXISTS idx_doc_lastmod ON documents(last_modified)')
            cursor.execute('CREATE INDEX IF NOT EXISTS idx_versions_docid ON document_versions(document_id)')
            cursor.execute('CREATE INDEX IF NOT EXISTS idx_detections_docid ON page_detections(document_id, page_number)')
            cursor.execute('CREATE INDEX IF NOT EXISTS idx_suggestions_field ON field_suggestions(field_name, frequency)')
            cursor.execute('CREATE INDEX IF NOT EXISTS idx_tags_docid ON document_tags(document_id)')

            # AI Search: bảng embeddings cho tìm kiếm ngữ nghĩa
            VectorStore.init_table(cursor)
            VectorStore.init_chunks_table(cursor)
            chunk_columns = {
                row[1] for row in cursor.execute("PRAGMA table_info(document_chunks)")
            }
            if "page_number" not in chunk_columns:
                cursor.execute(
                    "ALTER TABLE document_chunks ADD COLUMN page_number INTEGER"
                )

            # AI Search v2: FTS5 (BM25) — xem search_store.py
            import search_store
            search_store.init_fts(cursor)
            # Với external-content FTS5, count(*) trên bảng ảo phản chiếu
            # document_chunks kể cả khi index chưa có dữ liệu. Shadow table
            # docsize mới phản ánh số row thực sự đã được lập chỉ mục.
            fts_count = cursor.execute(
                "SELECT count(*) FROM chunks_fts_docsize"
            ).fetchone()[0]
            chunk_count = cursor.execute(
                "SELECT count(*) FROM document_chunks"
            ).fetchone()[0]
            if chunk_count and not fts_count:
                cursor.execute(
                    "INSERT INTO chunks_fts(chunks_fts) VALUES ('rebuild')"
                )

            # Legal knowledge graph: relations, entities, and inferred validity.
            import legal_graph
            legal_graph.init_tables(cursor)

            # Local assistant sessions and messages.
            import chat_history
            chat_history.init_tables(cursor)

            # AI Search v2: sqlite-vec KNN, tự fallback NumPy khi extension lỗi.
            # Rebuild whenever chunk_vectors does not mirror document_chunks
            # (ids, not just counts): databases written before chunk writes
            # kept the index in sync carry unindexed chunks and orphan vectors.
            import vector_index
            if vector_index.init_vec(conn):
                vector_index.VectorIndex(conn).ensure_synced(conn)

            conn.commit()

    def get_statistics(self):
        """Lấy thống kê từ database"""
        try:
            conn = self.conn_pool.get_connection()
            cursor = conn.cursor()
            stats = {}

            # Tổng số văn bản
            cursor.execute("SELECT COUNT(*) FROM documents")
            stats['total_documents'] = cursor.fetchone()[0]

            # Thống kê theo loại văn bản
            cursor.execute("""
                SELECT loai_vb, COUNT(*) as count 
                FROM document_versions 
                WHERE version_number = (
                    SELECT MAX(version_number) 
                    FROM document_versions v2 
                    WHERE v2.document_id = document_versions.document_id
                )
                GROUP BY loai_vb 
                ORDER BY count DESC
            """)
            stats['by_type'] = cursor.fetchall()

            # Thống kê theo độ khẩn
            cursor.execute("""
                SELECT do_khan, COUNT(*) as count 
                FROM document_versions 
                WHERE version_number = (
                    SELECT MAX(version_number) 
                    FROM document_versions v2 
                    WHERE v2.document_id = document_versions.document_id
                )
                GROUP BY do_khan 
                ORDER BY count DESC
            """)
            stats['by_urgency'] = cursor.fetchall()

            # Thống kê theo thời gian
            cursor.execute("""
                SELECT 
                    strftime('%Y-%m', created_at) as month,
                    COUNT(*) as count
                FROM documents
                GROUP BY month
                ORDER BY month DESC
                LIMIT 12
            """)
            stats['by_month'] = cursor.fetchall()

            # Văn bản mới nhất
            cursor.execute("""
                SELECT d.id, d.file_name, d.created_at, v.so_ki_hieu
                FROM documents d
                LEFT JOIN document_versions v ON d.id = v.document_id
                WHERE v.version_number = (
                    SELECT MAX(version_number) 
                    FROM document_versions 
                    WHERE document_id = d.id
                )
                ORDER BY d.created_at DESC
                LIMIT 5
            """)
            stats['recent_docs'] = cursor.fetchall()

            return stats

        except Exception as e:
            logger.error(f"Error getting statistics: {str(e)}")
            return None

    def create_backup(self, doc_id, file_path, reason="Manual backup"):
        """Create a backup of the document file"""
        try:
            if not os.path.exists(file_path):
                logger.warning(f"Cannot backup file that doesn't exist: {file_path}")
                return None

                
            backup_filename = f"{Path(file_path).stem}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{Path(file_path).suffix}"
            backup_path = BACKUP_DIR / backup_filename
            
            # Copy file to backup
            shutil.copy2(file_path, backup_path)
            
            # Register backup in database
            conn = self.conn_pool.get_connection()
            cursor = conn.cursor()
            cursor.execute(
                "INSERT INTO document_backups (document_id, backup_path, reason) VALUES (?, ?, ?)",
                (doc_id, str(backup_path), reason)
            )
            conn.commit()
                
            logger.info(f"Created backup of document {doc_id} at {backup_path}")
            return str(backup_path)
            
        except Exception as e:
            logger.error(f"Error creating document backup: {str(e)}")
            return None

    def load_suggestions(self):
        """Load suggestions from database into cache"""
        try:
            conn = self.conn_pool.get_connection()
            cursor = conn.cursor()
            cursor.execute('''
                SELECT field_name, value, frequency
                FROM field_suggestions
                ORDER BY frequency DESC, last_used DESC
            ''')
            
            for field, value, freq in cursor.fetchall():
                if field not in self.suggestions_cache:
                    self.suggestions_cache[field] = []
                self.suggestions_cache[field].append({
                    'value': value,
                    'frequency': freq
                })
                
        except Exception as e:
            logger.error(f"Error loading suggestions: {str(e)}")

    def get_suggestions(self, field_name: str, prefix: str = "") -> List[str]:
        """Get suggestions for a field, optionally filtered by prefix"""
        if field_name not in self.suggestions_cache:
            return []
            
        suggestions = self.suggestions_cache[field_name]
        if prefix:
            suggestions = [s for s in suggestions 
                         if s['value'].lower().startswith(prefix.lower())]
            
        return [s['value'] for s in sorted(
            suggestions, 
            key=lambda x: x['frequency'], 
            reverse=True
        )]

    def add_suggestion(self, field_name: str, value: str):
        """Add or update a suggestion in the database"""
        if not value or not value.strip():
            return
            
        value = value.strip()
            
        try:
            conn = self.conn_pool.get_connection()
            with conn:  # Auto commit/rollback
                cursor = conn.cursor()
                
                # Check if suggestion exists
                cursor.execute('''
                    SELECT id, frequency FROM field_suggestions
                    WHERE field_name = ? AND value = ?
                ''', (field_name, value))
                
                result = cursor.fetchone()
                if result:
                    cursor.execute('''
                        UPDATE field_suggestions
                        SET frequency = frequency + 1,
                            last_used = CURRENT_TIMESTAMP
                        WHERE id = ?
                    ''', (result[0],))
                else:
                    cursor.execute('''
                        INSERT INTO field_suggestions (field_name, value)
                        VALUES (?, ?)
                    ''', (field_name, value))
                    
            # Update cache
            if field_name not in self.suggestions_cache:
                self.suggestions_cache[field_name] = []
                
            # Find suggestion in cache
            found = False
            for suggestion in self.suggestions_cache[field_name]:

                if suggestion['value'] == value:
                    suggestion['frequency'] += 1
                    found = True
                    break
                    
            # Add to cache if not found
            if not found:
                self.suggestions_cache[field_name].append({
                    'value': value,
                    'frequency': 1
                })
                
        except Exception as e:
            logger.error(f"Error adding suggestion: {str(e)}")

    def add_document(self, file_path: str, ocr_results: Dict[str, Any], page_count: int = None, model_key: str = "chandra-q4") -> int:
        """Create a new document in the database"""
        try:
            if not os.path.exists(file_path):
                raise FileNotFoundError(f"File not found: {file_path}")
                
            # Calculate file hash for duplicate detection
            file_hash = self._calculate_file_hash(file_path)
            file_size = os.path.getsize(file_path)
            
            # Check if this document already exists
            existing_id = self._check_duplicate_document(file_path, file_hash)
            if existing_id:
                logger.info(f"Document {existing_id} exists, replacing OCR results with model {model_key}")
                # Tạo version mới với kết quả OCR mới thay vì bỏ qua
                self.replace_ocr_result(existing_id, ocr_results, model_key)
                return existing_id
            
            if page_count is None:
                ext = os.path.splitext(file_path)[1].lower()
                if ext in IMAGE_EXTENSIONS:
                    page_count = 1
                elif ext == '.pdf':
                    try:
                        from pdf2image.pdf2image import pdfinfo_from_path
                        info = pdfinfo_from_path(file_path, poppler_path=POPPLER_PATH)
                        page_count = info.get("Pages", 1)
                    except Exception:
                        page_count = 1
                    
            # Insert document với ID tiếp theo (luôn thêm vào cuối)
            conn = self.conn_pool.get_connection()
            with conn:
                cursor = conn.cursor()

                # Tìm ID lớn nhất hiện tại, văn bản mới sẽ có ID = max + 1
                cursor.execute('SELECT COALESCE(MAX(id), 0) + 1 FROM documents')
                result = cursor.fetchone()
                next_id = result[0] if result else 1

                cursor.execute('''
                    INSERT INTO documents (id, file_path, file_name, file_hash, file_size, page_count)
                    VALUES (?, ?, ?, ?, ?, ?)
                ''', (next_id, file_path, Path(file_path).name, file_hash, file_size, page_count))

                doc_id = next_id
                
                # Insert first version - bao gồm van_ban_dang_txt cho toàn bộ text OCR
                cursor.execute('''
                    INSERT INTO document_versions (
                        document_id, version_number, cqbh_tren, cqbh_duoi,
                        so_ki_hieu, loai_vb, nd_chinh, ngay_bh,
                        noi_nhan, chuc_vu, chu_ky, do_khan, do_mat, van_ban_dang_txt, modified_by, model_used
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ''', (
                    doc_id, 1,
                    ocr_results.get('CQBH_tren', ''),
                    ocr_results.get('CQBH_duoi', ''),
                    ocr_results.get('So_Ki_Hieu', ''),
                    ocr_results.get('Loai_VB', ''),
                    ocr_results.get('ND_Chinh', ''),
                    ocr_results.get('Ngay_BH', ''),
                    ocr_results.get('Noi_Nhan', ''),
                    ocr_results.get('Chuc_Vu', ''),
                    ocr_results.get('Chu_Ky', ''),
                    ocr_results.get('Do_Khan', 'Không'),
                    ocr_results.get('Do_Mat', ''),
                    ocr_results.get('Van_ban_dang_txt', ''),
                    'OCR System',
                    model_key
                ))

            
            # Add suggestions for all fields
            for field, value in {
                'so_ki_hieu': ocr_results.get('So_Ki_Hieu', ''),
                'loai_vb': ocr_results.get('Loai_VB', ''),
                'chuc_vu': ocr_results.get('Chuc_Vu', ''),
                'cqbh_tren': ocr_results.get('CQBH_tren', ''),
                'cqbh_duoi': ocr_results.get('CQBH_duoi', ''),
                'do_khan': ocr_results.get('Do_Khan', 'Không')
            }.items():
                if value:
                    self.add_suggestion(field, value)
            return doc_id
            
        except Exception as e:
            logger.error(f"Error adding document: {str(e)}")
            raise
    
    def replace_ocr_result(self, doc_id: int, ocr_results: Dict[str, Any],
                           model_key: str = "chandra-q4",
                           modified_by_suffix: str = ""):
        """Store a fresh OCR result of an existing document as a new version.

        Returns the new version number. The caller re-indexes the document
        (chunks, vectors, legal graph) afterwards.
        """
        updates = {
            'cqbh_tren': ocr_results.get('CQBH_tren', ''),
            'cqbh_duoi': ocr_results.get('CQBH_duoi', ''),
            'so_ki_hieu': ocr_results.get('So_Ki_Hieu', ''),
            'loai_vb': ocr_results.get('Loai_VB', ''),
            'nd_chinh': ocr_results.get('ND_Chinh', ''),
            'ngay_bh': ocr_results.get('Ngay_BH', ''),
            'noi_nhan': ocr_results.get('Noi_Nhan', ''),
            'chuc_vu': ocr_results.get('Chuc_Vu', ''),
            'chu_ky': ocr_results.get('Chu_Ky', ''),
            'do_khan': ocr_results.get('Do_Khan', 'Không'),
            'do_mat': ocr_results.get('Do_Mat', ''),
            'van_ban_dang_txt': ocr_results.get('Van_ban_dang_txt', ''),
        }
        # create_new_version reads model_used from the trailing "(...)".
        label = f"OCR {modified_by_suffix}".strip()
        return self.create_new_version(
            doc_id, updates, modified_by=f"{label} ({model_key})"
        )

    def _calculate_file_hash(self, file_path: str) -> str:
        """Calculate SHA-256 hash of a file for deduplication"""
        import hashlib
        
        try:
            sha256_hash = hashlib.sha256()
            with open(file_path, "rb") as f:
                # Read the file in chunks to handle large files efficiently
                for byte_block in iter(lambda: f.read(4096), b""):
                    sha256_hash.update(byte_block)
            return sha256_hash.hexdigest()
        except Exception as e:
            logger.error(f"Error calculating file hash: {str(e)}")
            return ""
            
    def _check_duplicate_document(self, file_path: str, file_hash: str) -> Optional[int]:
        """Check if document already exists by path or hash"""
        if not file_hash:
            return None
            
        try:
            conn = self.conn_pool.get_connection()
            cursor = conn.cursor()
            
            # First check by file path
            cursor.execute(
                "SELECT id FROM documents WHERE file_path = ? LIMIT 1",
                (file_path,)
            )
            result = cursor.fetchone()
            if result:
                return result[0]
                
            # Then check by hash if provided
            cursor.execute(
                "SELECT id FROM documents WHERE file_hash = ? LIMIT 1",
                (file_hash,)
            )
            result = cursor.fetchone()
            if result:
                return result[0]
                
            return None
        except Exception as e:
            logger.error(f"Error checking for duplicate document: {str(e)}")
            return None

    def add_page_detections(self, doc_id: int, page_number: int, detections: List[Dict], page_text: str = None):
        """Save detections for a specific page"""
        try:
            if isinstance(detections, str):
                detections = json.loads(detections)
            conn = self.conn_pool.get_connection()
            cursor = conn.cursor()
            
            # Check if detections already exist for this page
            cursor.execute(
                "SELECT id FROM page_detections WHERE document_id = ? AND page_number = ?",
                (doc_id, page_number)
            )
            
            existing = cursor.fetchone()
            if existing:
                cursor.execute(
                    "UPDATE page_detections SET detection_data = ?, page_text = ? WHERE id = ?",
                    (json.dumps(detections), page_text, existing[0])
                )
            else:
                cursor.execute('''
                    INSERT INTO page_detections (document_id, page_number, detection_data, page_text)
                    VALUES (?, ?, ?, ?)
                ''', (doc_id, page_number, json.dumps(detections), page_text))
                
            conn.commit()
        except Exception as e:
            logger.error(f"Error adding page detections: {str(e)}")
            raise

    def get_document_detections(self, doc_id: int, page_number: int = None) -> Union[List[Dict], Dict[int, List[Dict]]]:
        """Get detections for a document page or all pages"""
        try:
            conn = self.conn_pool.get_connection()

            cursor = conn.cursor()
            
            if page_number is not None:
                cursor.execute('''
                    SELECT detection_data
                    FROM page_detections
                    WHERE document_id = ? AND page_number = ?
                ''', (doc_id, page_number))
                
                result = cursor.fetchone()
                if result and result[0]:
                    try:
                        # Set higher recursion limit temporarily
                        current_limit = sys.getrecursionlimit()
                        sys.setrecursionlimit(10000)
                        decoded = json.loads(result[0])
                        sys.setrecursionlimit(current_limit)
                        return decoded
                    except json.JSONDecodeError:
                        return []
                return []
            else:
                cursor.execute('''
                    SELECT page_number, detection_data
                    FROM page_detections
                    WHERE document_id = ?
                    ORDER BY page_number
                ''', (doc_id,))
                
                results = {}
                for page_num, detection_data in cursor.fetchall():
                    if detection_data:
                        try:
                            results[page_num] = json.loads(detection_data)
                        except json.JSONDecodeError:
                            results[page_num] = []
                return results
                
        except Exception as e:
            print(f"Error getting document detections: {str(e)}")  # Use print instead of logger
            return [] if page_number is not None else {}

    def delete_document(self, doc_id: int, keep_file: bool = False):
        """Delete document and all related data"""
        try:
            conn = self.conn_pool.get_connection()
            with conn:
                cursor = conn.cursor()
                
                # Backup file path before deleting
                cursor.execute('SELECT file_path FROM documents WHERE id = ?', (doc_id,))
                file_path_result = cursor.fetchone()
                file_path = file_path_result[0] if file_path_result else None
                chunk_ids = [
                    row[0] for row in cursor.execute(
                        'SELECT id FROM document_chunks WHERE document_id = ?',
                        (doc_id,),
                    )
                ]
                affected_validity = [
                    row[0] for row in cursor.execute(
                        'SELECT DISTINCT target_doc_id FROM doc_relations '
                        'WHERE source_doc_id = ? AND target_doc_id IS NOT NULL '
                        'AND target_doc_id != ?',
                        (doc_id, doc_id),
                    )
                ]
                
                # Create backup if file exists
                if file_path and os.path.exists(file_path):
                    self.create_backup(doc_id, file_path, reason="Pre-deletion backup")
                
                # Delete related records in correct order
                # 0. Delete AI search embedding + chunks
                from legal_graph import GraphStore, compute_validity
                from vector_index import VectorIndex

                GraphStore(conn).delete_doc(doc_id)
                VectorIndex(conn).remove_chunks(conn, chunk_ids)
                cursor.execute('DELETE FROM document_embeddings WHERE document_id = ?', (doc_id,))
                VectorStore.delete_chunks(cursor, doc_id)

                # 1. Delete page detections
                cursor.execute('DELETE FROM page_detections WHERE document_id = ?', (doc_id,))
                
                # 2. Delete document tags
                cursor.execute('DELETE FROM document_tags WHERE document_id = ?', (doc_id,))
                
                # 3. Delete document versions
                cursor.execute('DELETE FROM document_versions WHERE document_id = ?', (doc_id,))
                
                # 4. Delete from document backups
                cursor.execute('DELETE FROM document_backups WHERE document_id = ?', (doc_id,))
                
                # 5. Finally delete the document
                cursor.execute('DELETE FROM documents WHERE id = ?', (doc_id,))

                if affected_validity:
                    compute_validity(conn, affected_validity)

                # Commit the transaction
                conn.commit()
                
                # Delete file if it exists and not keeping
                if not keep_file and file_path and os.path.exists(file_path):
                    try:
                        os.remove(file_path)
                        logger.info(f"Deleted file: {file_path}")
                    except Exception as e:
                        logger.warning(f"Cannot delete file: {file_path}. Error: {str(e)}")
                        
                return True
                
        except Exception as e:
            logger.error(f"Error deleting document: {str(e)}")
            if conn:
                conn.rollback()
            raise

    def _parse_vietnamese_date(self, date_string):
        """
        Parse Vietnamese date string to standard format

        """
        if not isinstance(date_string, str):
            return date_string
            
        try:
            # Chuẩn hóa chuỗi đầu vào
            date_string = date_string.lower().strip()
            
            # Xử lý format chuẩn từ database (yyyy-mm-dd)
            if '-' in date_string and len(date_string.split('-')[0]) == 4:
                return date_string
                
            # Xử lý format "ngày dd tháng mm năm yyyy"
            if "ngày" in date_string and "tháng" in date_string and "năm" in date_string:
                date_string = date_string.replace("ngày", "").replace("tháng", "").replace("năm", "")
                parts = [x.strip() for x in date_string.split() if x.strip()]
                if len(parts) >= 3:
                    day = int(parts[0])
                    month = int(parts[1])
                    year = int(parts[2])
                    return f"{day:02d}/{month:02d}/{year:04d}"
                    
            # Xử lý format dd/mm/yyyy
            if "/" in date_string:
                parts = date_string.split("/")
                if len(parts) == 3:
                    day = int(parts[0])
                    month = int(parts[1])
                    year = int(parts[2])
                    return f"{day:02d}/{month:02d}/{year:04d}"
                    
            return date_string
            
        except Exception as e:
            logger.error(f"Error parsing date string '{date_string}': {str(e)}")
            return date_string

    def export_to_json(self, output_path: str, filter_criteria: Dict = None):
        """Export database to JSON format"""
        try:
            conn = self.conn_pool.get_connection()
            # Base query với các cột được sắp xếp hợp lý
            query = '''
                SELECT 
                    d.id as "ID",
                    d.file_name as "Tên File",
                    d.created_at as "Ngày Tạo",
                    v.cqbh_tren as "CQBH Trên",
                    v.cqbh_duoi as "CQBH Dưới",
                    v.so_ki_hieu as "Số Ký Hiệu",
                    v.ngay_bh as "Ngày Ban Hành",
                    v.do_khan as "Độ Khẩn",
                    v.do_mat as "Độ Mật",
                    v.loai_vb as "Loại Văn Bản",
                    v.nd_chinh as "Nội Dung Chính",
                    v.noi_nhan as "Nơi Nhận",
                    v.chuc_vu as "Chức Vụ",
                    v.chu_ky as "Chữ Ký",
                    v.van_ban_dang_txt as "Nội Dung Văn Bản (OCR)"
                FROM documents d
                LEFT JOIN document_versions v ON d.id = v.document_id
                WHERE v.version_number = (
                    SELECT MAX(version_number) 
                    FROM document_versions 
                    WHERE document_id = d.id
                )
            '''

            # Xử lý filter criteria nếu có
            params = []
            if filter_criteria:
                conditions = []
                for field, value in filter_criteria.items():
                    if value:
                        if field == 'id':
                            conditions.append("d.id = ?")
                            params.append(value)
                        elif field == 'file_name':
                            conditions.append("d.file_name LIKE ?")
                            params.append(f'%{value}%')
                        elif field == 'date_from':
                            conditions.append("d.created_at >= ?")
                            params.append(value)
                        elif field == 'date_to':
                            conditions.append("d.created_at <= ?")
                            params.append(value)
                        elif field in ['cqbh_tren', 'cqbh_duoi', 'so_ki_hieu', 'loai_vb',
                                    'do_khan', 'do_mat', 'ngay_bh', 'chuc_vu']:
                            conditions.append(f"v.{field} LIKE ?")
                            params.append(f'%{value}%')
                        elif field == 'nd_chinh':
                            conditions.append("v.nd_chinh LIKE ?")
                            params.append(f'%{value}%')
                
                if conditions:
                    query += " AND " + " AND ".join(conditions)
            
            query += " ORDER BY d.created_at DESC"
            
            # Thực thi query

            if params:
                df = pd.read_sql_query(query, conn, params=params)
            else:
                df = pd.read_sql_query(query, conn)
            
            # Xử lý datetime columns
            datetime_columns = ['Ngày Tạo', 'Ngày Ban Hành']
            for col in datetime_columns:
                if col in df.columns:
                    # Áp dụng hàm parse cho từng giá trị trong cột
                    df[col] = df[col].apply(self._parse_vietnamese_date)

            # Chuyển đổi DataFrame thành JSON
            result_json = df.to_json(orient='records', force_ascii=False, date_format='iso')
            
            # Ghi ra file với encoding UTF-8
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(result_json)

            logger.info(f"Successfully exported data to JSON: {output_path}")
            return True
                
        except Exception as e:
            logger.error(f"Error exporting to JSON: {str(e)}")
            raise
            
    def export_to_excel(self, output_path: str, filter_criteria: Dict = None):
        """Export database to Excel with optimized formatting"""
        try:
            # Kiểm tra thư viện
            try:
                import pandas as pd
                from openpyxl import styles
                from openpyxl.utils import get_column_letter
            except ImportError:
                raise ImportError("Thư viện 'openpyxl' chưa được cài đặt.")
                    
            conn = self.conn_pool.get_connection()
            # Base query với các cột được sắp xếp hợp lý
            query = '''
                SELECT 
                    d.id as "ID",
                    d.file_name as "Tên File",
                    d.created_at as "Ngày Tạo",
                    v.cqbh_tren as "CQBH Trên",
                    v.cqbh_duoi as "CQBH Dưới",
                    v.so_ki_hieu as "Số Ký Hiệu",
                    v.ngay_bh as "Ngày Ban Hành",
                    v.do_khan as "Độ Khẩn",
                    v.do_mat as "Độ Mật",
                    v.loai_vb as "Loại Văn Bản",
                    v.nd_chinh as "Nội Dung Chính",
                    v.noi_nhan as "Nơi Nhận",
                    v.chuc_vu as "Chức Vụ",
                    v.chu_ky as "Chữ Ký",
                    v.van_ban_dang_txt as "Nội Dung Văn Bản (OCR)"
                FROM documents d
                LEFT JOIN document_versions v ON d.id = v.document_id
                WHERE v.version_number = (
                    SELECT MAX(version_number) 
                    FROM document_versions 
                    WHERE document_id = d.id
                )
            '''

            # Xử lý filter criteria nếu có
            params = []
            if filter_criteria:
                conditions = []
                for field, value in filter_criteria.items():
                    if value:
                        if field == 'id':
                            conditions.append("d.id = ?")
                            params.append(value)
                        elif field == 'file_name':
                            conditions.append("d.file_name LIKE ?")
                            params.append(f'%{value}%')
                        elif field == 'date_from':
                            conditions.append("d.created_at >= ?")
                            params.append(value)
                        elif field == 'date_to':
                            conditions.append("d.created_at <= ?")
                            params.append(value)
                        elif field in ['cqbh_tren', 'cqbh_duoi', 'so_ki_hieu', 'loai_vb',
                                    'do_khan', 'do_mat', 'ngay_bh', 'chuc_vu']:
                            conditions.append(f"v.{field} LIKE ?")
                            params.append(f'%{value}%')
                        elif field == 'nd_chinh':
                            conditions.append("v.nd_chinh LIKE ?")
                            params.append(f'%{value}%')
                
                if conditions:
                    query += " AND " + " AND ".join(conditions)
            
            query += " ORDER BY d.created_at DESC"
            
            # Thực thi query
            if params:
                df = pd.read_sql_query(query, conn, params=params)
            else:

                df = pd.read_sql_query(query, conn)
            
            # Xử lý datetime columns
            datetime_columns = ['Ngày Tạo', 'Ngày Ban Hành']
            for col in datetime_columns:
                if col in df.columns:
                    df[col] = df[col].apply(self._parse_vietnamese_date)

            # Strip HTML tags from van_ban_dang_txt (Chandra outputs raw HTML)
            ocr_col = "Nội Dung Văn Bản (OCR)"
            if ocr_col in df.columns:
                from bs4 import BeautifulSoup
                def _strip_html(val):
                    if not val or not isinstance(val, str):
                        return val
                    if '<div' in val or '<p>' in val:
                        return BeautifulSoup(val, 'html.parser').get_text(separator='\n', strip=True)
                    return val
                df[ocr_col] = df[ocr_col].apply(_strip_html)

            # Export to Excel với formatting tối ưu
            with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
                df.to_excel(writer, index=False, sheet_name='Documents')
                
                workbook = writer.book
                worksheet = writer.sheets['Documents']
                
                # Định nghĩa styles
                header_style = styles.NamedStyle(name='header_style')
                header_style.font = styles.Font(bold=True, size=11)
                header_style.fill = styles.PatternFill(start_color='E0E0E0', end_color='E0E0E0', fill_type='solid')
                header_style.alignment = styles.Alignment(horizontal='center', vertical='center', wrap_text=True)
                header_style.border = styles.Border(
                    left=styles.Side(style='thin'),
                    right=styles.Side(style='thin'),
                    top=styles.Side(style='thin'),
                    bottom=styles.Side(style='thin')
                )

                # Style cho dữ liệu
                data_style = styles.NamedStyle(name='data_style')
                data_style.font = styles.Font(size=10)
                data_style.alignment = styles.Alignment(vertical='center', wrap_text=True)
                data_style.border = styles.Border(
                    left=styles.Side(style='thin'),
                    right=styles.Side(style='thin'),
                    top=styles.Side(style='thin'),
                    bottom=styles.Side(style='thin')
                )

                # Cấu hình độ rộng và style cho từng cột
                column_widths = {
                    'ID': 8,
                    'Tên File': 25,
                    'Ngày Tạo': 12,
                    'CQBH Trên': 25,
                    'CQBH Dưới': 25,
                    'Số Ký Hiệu': 20,
                    'Loại Văn Bản': 15,
                    'Nội Dung Chính': 40,
                    'Ngày Ban Hành': 15,
                    'Nơi Nhận': 50,
                    'Chức Vụ': 20,
                    'Chữ Ký': 20,
                    'Độ Khẩn': 12
                }

                # Áp dụng style và độ rộng cho các cột
                for idx, column in enumerate(df.columns, 1):
                    col_letter = get_column_letter(idx)
                    
                    # Áp dụng style cho header
                    cell = worksheet.cell(row=1, column=idx)
                    cell.style = header_style
                    
                    # Set độ rộng cột
                    width = column_widths.get(column, 15)  # Default width là 15 nếu không được định nghĩa
                    worksheet.column_dimensions[col_letter].width = width
                    
                    # Áp dụng style cho tất cả cells trong cột
                    for row in range(2, worksheet.max_row + 1):
                        cell = worksheet.cell(row=row, column=idx)
                        cell.style = data_style
                        
                        # Căn giữa cho một số cột cụ thể
                        if column in ['ID', 'Ngày Tạo', 'Ngày Ban Hành', 'Độ Khẩn']:
                            cell.alignment = styles.Alignment(horizontal='center', vertical='center', wrap_text=True)
                        
                        # Căn trái cho các cột còn lại
                        else:
                            cell.alignment = styles.Alignment(horizontal='left', vertical='center', wrap_text=True)

                # Set độ cao cho header
                worksheet.row_dimensions[1].height = 35

                # Set độ cao cho data rows
                for row in range(2, worksheet.max_row + 1):
                    # Tính toán độ cao dựa trên nội dung
                    max_length = 0
                    for cell in worksheet[row]:

                        if cell.value:
                            lines = str(cell.value).count('\n') + 1
                            max_length = max(max_length, lines)
                    
                    # Set độ cao tối thiểu 20, và thêm 15 cho mỗi dòng nếu có nhiều dòng
                    row_height = max(20, min(15 * max_length, 100))  # giới hạn độ cao tối đa là 100
                    worksheet.row_dimensions[row].height = row_height

                # Freeze panes
                worksheet.freeze_panes = 'A2'
                
                # Auto-filter
                worksheet.auto_filter.ref = worksheet.dimensions

            logger.info(f"Successfully exported data to {output_path}")
            return True
                
        except Exception as e:
            logger.error(f"Error exporting to Excel: {str(e)}")
            raise

    def get_all_documents(self, filter_criteria: Dict = None, sort_by: str = 'created_at', sort_desc: bool = True):
        """Get all documents with optional filtering and sorting"""
        try:
            conn = self.conn_pool.get_connection()
            query = '''
                SELECT 
                    d.id, 
                    d.file_path,
                    d.file_name,
                    d.created_at,
                    d.page_count,
                    (SELECT COUNT(v.id) FROM document_versions v WHERE v.document_id = d.id) as version_count,
                    (SELECT cqbh_tren FROM document_versions 
                     WHERE document_id = d.id 
                     ORDER BY version_number DESC LIMIT 1) as cqbh_tren,
                    (SELECT cqbh_duoi FROM document_versions 
                     WHERE document_id = d.id 
                     ORDER BY version_number DESC LIMIT 1) as cqbh_duoi,
                    (SELECT so_ki_hieu FROM document_versions 
                     WHERE document_id = d.id 
                     ORDER BY version_number DESC LIMIT 1) as so_ki_hieu,
                    (SELECT loai_vb FROM document_versions 
                     WHERE document_id = d.id 
                     ORDER BY version_number DESC LIMIT 1) as loai_vb,
                    (SELECT do_khan FROM document_versions
                     WHERE document_id = d.id
                     ORDER BY version_number DESC LIMIT 1) as do_khan,
                    (SELECT modified_by FROM document_versions
                     WHERE document_id = d.id
                     ORDER BY version_number DESC LIMIT 1) as modified_by,
                    (SELECT created_at FROM document_versions
                     WHERE document_id = d.id
                     ORDER BY version_number DESC LIMIT 1) as last_modified,
                    (SELECT model_used FROM document_versions
                     WHERE document_id = d.id
                     ORDER BY version_number DESC LIMIT 1) as model_used
                FROM documents d
            '''

            # Add filter conditions if provided
            params = []
            if filter_criteria:
                conditions = []
                for field, value in filter_criteria.items():
                    if value:
                        if field == 'file_name':
                            conditions.append(f"d.file_name LIKE ?")
                            params.append(f'%{value}%')
                        elif field == 'text':
                            text_condition = '''
                                EXISTS (
                                    SELECT 1 FROM document_versions v 
                                    WHERE v.document_id = d.id 
                                    AND (
                                        v.cqbh_tren LIKE ? OR
                                        v.cqbh_duoi LIKE ? OR
                                        v.so_ki_hieu LIKE ? OR
                                        v.loai_vb LIKE ? OR
                                        v.nd_chinh LIKE ? OR
                                        v.ngay_bh LIKE ? OR
                                        v.noi_nhan LIKE ? OR 
                                        v.chuc_vu LIKE ? OR
                                        v.chu_ky LIKE ? OR
                                        v.do_khan LIKE ? OR
                                        v.van_ban_dang_txt LIKE ?
                                    )
                                )
                            '''
                            conditions.append(text_condition)
                            params.extend([f'%{value}%'] * 11)  # One for each field including van_ban_dang_txt
                        elif field == 'date_from':
                            conditions.append(f"d.created_at >= ?")
                            params.append(value)
                        elif field == 'date_to':
                            conditions.append(f"d.created_at <= ?")
                            params.append(value)
                        elif field == 'do_khan':
                            conditions.append('''
                                EXISTS (

                                    SELECT 1 FROM document_versions v
                                    WHERE v.document_id = d.id
                                    AND v.version_number = (
                                        SELECT MAX(version_number) FROM document_versions
                                        WHERE document_id = d.id
                                    )
                                    AND v.do_khan = ?
                                )
                            ''')
                            params.append(value)
                
                if conditions:
                    query += " WHERE " + " AND ".join(conditions)
            
            # Add sorting
            sort_column = sort_by
            if sort_by == 'so_ki_hieu':
                sort_column = "(SELECT so_ki_hieu FROM document_versions WHERE document_id = d.id ORDER BY version_number DESC LIMIT 1)"
            elif sort_by == 'do_khan':
                sort_column = "(SELECT do_khan FROM document_versions WHERE document_id = d.id ORDER BY version_number DESC LIMIT 1)"
                
            query += f" ORDER BY {sort_column} {'DESC' if sort_desc else 'ASC'}"
            
            # Execute query
            cursor = conn.cursor()
            if params:
                cursor.execute(query, params)
            else:
                cursor.execute(query)
                
            return cursor.fetchall()
                
        except Exception as e:
            logger.error(f"Error fetching documents: {str(e)}")
            return []

    def get_latest_previews_batch(self, doc_ids: list) -> dict:
        """Batched tooltip preview fetch — 1 query for N docs.

        Performance fix 2026-04-18: previous `_rebuild_doc_list_only()`
        called `get_latest_version(doc_id)` in a loop, causing N+1 queries
        (300 docs × 5ms = 1.5s UI freeze). This method issues ONE query
        with IN clause and SUBSTR to fetch only 2KB of van_ban per row
        instead of 75KB.

        Returns:
            dict[doc_id → {loai_vb, nd_chinh, van_ban_preview}]
        """
        if not doc_ids:
            return {}
        try:
            conn = self.conn_pool.get_connection()
            cursor = conn.cursor()
            # Subquery picks MAX version_number per document_id.
            placeholders = ",".join("?" * len(doc_ids))
            query = f"""
                SELECT dv.document_id,
                       dv.loai_vb,
                       dv.nd_chinh,
                       SUBSTR(COALESCE(dv.van_ban_dang_txt, ''), 1, 2000) AS vb_preview
                FROM document_versions dv
                INNER JOIN (
                    SELECT document_id, MAX(version_number) AS max_ver
                    FROM document_versions
                    WHERE document_id IN ({placeholders})
                    GROUP BY document_id
                ) latest ON dv.document_id = latest.document_id
                       AND dv.version_number = latest.max_ver
            """
            cursor.execute(query, doc_ids)
            result = {}
            for row in cursor.fetchall():
                result[row[0]] = {
                    "loai_vb": (row[1] or "").strip(),
                    "nd_chinh": (row[2] or "").strip(),
                    "van_ban_preview": (row[3] or "").strip(),
                }
            return result
        except Exception as e:
            logger.error(f"Error in batch preview fetch: {e}")
            return {}

    def get_latest_version(self, doc_id):
        """Get the most recent version of a document"""
        try:
            conn = self.conn_pool.get_connection()
            cursor = conn.cursor()
            cursor.execute('''
                SELECT * FROM document_versions
                WHERE document_id = ?
                ORDER BY version_number DESC
                LIMIT 1
            ''', (doc_id,))
            return cursor.fetchone()
        except Exception as e:
            logger.error(f"Error getting latest version: {str(e)}")
            return None
        
    def get_document_versions(self, doc_id):
        """Get all versions of a document"""

        try:
            conn = self.conn_pool.get_connection()
            cursor = conn.cursor()
            cursor.execute('''
                SELECT * FROM document_versions
                WHERE document_id = ?
                ORDER BY version_number DESC
            ''', (doc_id,))
            return cursor.fetchall()
        except Exception as e:
            logger.error(f"Error getting document versions: {str(e)}")
            return []
        
    def get_document_info(self, doc_id):
        """Get basic document information"""
        try:
            conn = self.conn_pool.get_connection()
            cursor = conn.cursor()
            cursor.execute('''
                SELECT * FROM documents
                WHERE id = ?
            ''', (doc_id,))
            return cursor.fetchone()
        except Exception as e:
            logger.error(f"Error getting document info: {str(e)}")
            return None

    def create_new_version(self, doc_id: int, updates: Dict[str, str], modified_by: str = "User"):
        """Create a new version of a document"""
        try:
            # Đầu tiên kiểm tra xem document có tồn tại không
            conn = self.conn_pool.get_connection()
            cursor = conn.cursor()
            cursor.execute('SELECT id FROM documents WHERE id = ?', (doc_id,))
            if not cursor.fetchone():
                raise ValueError(f"Document with ID {doc_id} does not exist")

            with conn:
                # Get current version number
                cursor.execute('''
                    SELECT MAX(version_number)
                    FROM document_versions
                    WHERE document_id = ?
                ''', (doc_id,))
                
                current_version = cursor.fetchone()[0] or 0
                new_version = current_version + 1

                # Xác định model_used từ modified_by nếu có
                model_used = None
                if 'OCR' in modified_by and '(' in modified_by:
                    model_used = modified_by.split('(')[-1].rstrip(')')
                if model_used is None:
                    cursor.execute('''
                        SELECT model_used FROM document_versions
                        WHERE document_id = ?
                        ORDER BY version_number DESC LIMIT 1
                    ''', (doc_id,))
                    previous = cursor.fetchone()
                    model_used = previous[0] if previous else ACTIVE_OCR_MODEL

                # Chuẩn bị dữ liệu cho version mới
                version_data = (
                    doc_id,
                    new_version,
                    updates.get('cqbh_tren', ''),
                    updates.get('cqbh_duoi', ''),
                    updates.get('so_ki_hieu', ''),
                    updates.get('loai_vb', ''),
                    updates.get('nd_chinh', ''),
                    updates.get('ngay_bh', ''),
                    updates.get('noi_nhan', ''),
                    updates.get('chuc_vu', ''),
                    updates.get('chu_ky', ''),
                    updates.get('do_khan', 'Không'),
                    updates.get('do_mat', ''),
                    updates.get('van_ban_dang_txt', ''),  # Toàn bộ text OCR
                    modified_by,
                    model_used
                )

                # Insert new version with error handling
                try:
                    # Cập nhật lại câu lệnh SQL để lưu TIMESTAMP rõ ràng
                    cursor.execute('''
                        INSERT INTO document_versions (
                            document_id, version_number, cqbh_tren, cqbh_duoi,
                            so_ki_hieu, loai_vb, nd_chinh, ngay_bh,
                            noi_nhan, chuc_vu, chu_ky, do_khan, do_mat, van_ban_dang_txt, modified_by, model_used
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    ''', version_data)

                    # Update last_modified in documents table
                    cursor.execute('''
                        UPDATE documents
                        SET last_modified = CURRENT_TIMESTAMP
                        WHERE id = ?
                    ''', (doc_id,))

                    conn.commit()

                    
                    # Add suggestions for fields
                    for field in ['so_ki_hieu', 'loai_vb', 'chuc_vu', 'cqbh_tren', 'cqbh_duoi', 'do_khan', 'do_mat']:
                        if field in updates and updates[field].strip():
                            self.add_suggestion(field, updates[field].strip())
                    
                    return new_version

                except sqlite3.IntegrityError as e:
                    conn.rollback()
                    logger.error(f"Database integrity error: {str(e)}")
                    raise
                except Exception as e:
                    conn.rollback()
                    logger.error(f"Error creating new version: {str(e)}")
                    raise
                    
        except Exception as e:
            logger.error(f"Error creating new version: {str(e)}")
            raise

    def search_documents(self, query: str, search_type: str = "all"):
        """
        Advanced search with multi-keyword support, frequency counting, and snippet extraction.
        
        Keywords are separated by ';'. Results are sorted by total keyword frequency (descending).
        
        Returns:
            List of tuples: (document_data, search_stats)
            where search_stats = {
                'total_frequency': int,
                'keyword_frequencies': {keyword: count},
                'snippets': [{'keyword': str, 'line': int, 'context': str}, ...]
            }
        """
        if not query:
            # Return all documents with empty stats
            docs = self.get_all_documents()
            return [(doc, {'total_frequency': 0, 'keyword_frequencies': {}, 'snippets': []}) for doc in docs]
        
        # Parse keywords (split by ';' and strip whitespace)
        keywords = [kw.strip() for kw in query.split(';') if kw.strip()]
        if not keywords:
            docs = self.get_all_documents()
            return [(doc, {'total_frequency': 0, 'keyword_frequencies': {}, 'snippets': []}) for doc in docs]
        
        # Build filter criteria to find candidate documents
        filter_criteria = {}
        if search_type == "file_name":
            filter_criteria['file_name'] = keywords[0]  # Use first keyword for file name search
        elif search_type == "do_khan":
            filter_criteria['do_khan'] = keywords[0]
        else:
            # For content/all/so_ki_hieu, use text search with first keyword
            filter_criteria['text'] = keywords[0]
        
        # Get candidate documents
        documents = self.get_all_documents(filter_criteria)
        
        # For multi-keyword, we need to also include documents matching other keywords
        if len(keywords) > 1 and search_type not in ["file_name", "do_khan"]:
            doc_ids = set(doc[0] for doc in documents)
            for kw in keywords[1:]:
                extra_docs = self.get_all_documents({'text': kw})
                for doc in extra_docs:
                    if doc[0] not in doc_ids:
                        documents.append(doc)
                        doc_ids.add(doc[0])
        
        # Additional filtering for so_ki_hieu
        if search_type == "so_ki_hieu":
            documents = [doc for doc in documents if any(kw.lower() in (doc[8] or '').lower() for kw in keywords)]
        
        # Now analyze each document for frequency and snippets
        results = []
        for doc in documents:
            doc_id = doc[0]
            stats = self._analyze_document_keywords(doc_id, keywords)
            results.append((doc, stats))
        
        # Sort by total_frequency descending
        results.sort(key=lambda x: x[1]['total_frequency'], reverse=True)
        
        return results
    
    def _analyze_document_keywords(self, doc_id: int, keywords: list) -> dict:
        """
        Analyze a document for keyword frequencies and extract snippets.
        
        Args:
            doc_id: Document ID
            keywords: List of keywords to search for
            
        Returns:
            dict with 'total_frequency', 'keyword_frequencies', 'snippets'
        """
        try:
            conn = self.conn_pool.get_connection()
            cursor = conn.cursor()
            

            # Get the latest version's text content
            cursor.execute('''
                SELECT van_ban_dang_txt, nd_chinh, so_ki_hieu, cqbh_tren, cqbh_duoi, 
                       loai_vb, noi_nhan, chuc_vu, chu_ky
                FROM document_versions 
                WHERE document_id = ? 
                ORDER BY version_number DESC LIMIT 1
            ''', (doc_id,))
            
            row = cursor.fetchone()
            if not row:
                return {'total_frequency': 0, 'keyword_frequencies': {}, 'snippets': []}
            
            # Combine all text fields for searching
            text_fields = [row[i] or '' for i in range(len(row))]
            full_text = '\n'.join(text_fields)
            
            # Also get page_text from page_detections
            cursor.execute('''
                SELECT page_number, page_text FROM page_detections 
                WHERE document_id = ? ORDER BY page_number
            ''', (doc_id,))
            page_rows = cursor.fetchall()
            for page_row in page_rows:
                if page_row[1]:
                    full_text += f'\n--- Trang {page_row[0]} ---\n' + page_row[1]
            
            # Analyze keywords
            keyword_frequencies = {}
            snippets = []
            full_text_lower = full_text.lower()
            lines = full_text.split('\n')
            
            for kw in keywords:
                kw_lower = kw.lower()
                # Count frequency
                count = full_text_lower.count(kw_lower)
                keyword_frequencies[kw] = count
                
                # Extract snippets (up to 3 per keyword)
                snippet_count = 0
                for line_num, line in enumerate(lines, 1):
                    if kw_lower in line.lower() and snippet_count < 3:
                        # Get context (the line itself, trimmed)
                        context = line.strip()
                        if len(context) > 150:
                            # Find keyword position and show surrounding context
                            pos = line.lower().find(kw_lower)
                            start = max(0, pos - 50)
                            end = min(len(line), pos + len(kw) + 50)
                            context = ('...' if start > 0 else '') + line[start:end].strip() + ('...' if end < len(line) else '')
                        
                        snippets.append({
                            'keyword': kw,
                            'line': line_num,
                            'context': context
                        })
                        snippet_count += 1
            
            total_frequency = sum(keyword_frequencies.values())
            
            return {
                'total_frequency': total_frequency,
                'keyword_frequencies': keyword_frequencies,
                'snippets': snippets
            }
            
        except Exception as e:
            logger.error(f"Error analyzing document {doc_id} for keywords: {str(e)}")
            return {'total_frequency': 0, 'keyword_frequencies': {}, 'snippets': []}
        
    def verify_file_paths(self):
        """Check and identify missing file paths"""
        try:
            conn = self.conn_pool.get_connection()
            cursor = conn.cursor()
            cursor.execute('SELECT id, file_path FROM documents')
            documents = cursor.fetchall()
            
            missing_files = []
            for doc_id, file_path in documents:
                if not os.path.exists(file_path):
                    missing_files.append((doc_id, file_path))
                    
            return missing_files
        except Exception as e:
            logger.error(f"Error verifying file paths: {str(e)}")
            return []

    def update_file_path(self, doc_id: int, new_path: str):
        """Update file path for a document"""
        try:
            if not os.path.exists(new_path):
                raise FileNotFoundError(f"New file path does not exist: {new_path}")
                
            # Calculate new hash and size
            file_hash = self._calculate_file_hash(new_path)
            file_size = os.path.getsize(new_path)
            
            # Update page count if PDF

            page_count = None
            if Path(new_path).suffix.lower() == '.pdf':
                try:
                    with fitz.open(new_path) as doc:
                        page_count = len(doc)
                except:
                    pass
                    
            conn = self.conn_pool.get_connection()
            with conn:
                cursor = conn.cursor()
                if page_count is not None:
                    cursor.execute(
                        'UPDATE documents SET file_path = ?, file_hash = ?, file_size = ?, page_count = ? WHERE id = ?',
                        (new_path, file_hash, file_size, page_count, doc_id)
                    )
                else:
                    cursor.execute(
                        'UPDATE documents SET file_path = ?, file_hash = ?, file_size = ? WHERE id = ?',
                        (new_path, file_hash, file_size, doc_id)
                    )
                
                return cursor.rowcount > 0
        except Exception as e:
            logger.error(f"Error updating file path: {str(e)}")
            raise

    def add_tag(self, doc_id: int, tag_name: str):
        """Add a tag to a document"""
        try:
            conn = self.conn_pool.get_connection()
            with conn:
                cursor = conn.cursor()
                cursor.execute(
                    'INSERT OR IGNORE INTO document_tags (document_id, tag_name) VALUES (?, ?)',
                    (doc_id, tag_name)
                )
                return cursor.rowcount > 0
        except Exception as e:
            logger.error(f"Error adding tag: {str(e)}")
            return False
            
    def remove_tag(self, doc_id: int, tag_name: str):
        """Remove a tag from a document"""
        try:
            conn = self.conn_pool.get_connection()
            with conn:
                cursor = conn.cursor()
                cursor.execute(
                    'DELETE FROM document_tags WHERE document_id = ? AND tag_name = ?',
                    (doc_id, tag_name)
                )
                return cursor.rowcount > 0
        except Exception as e:
            logger.error(f"Error removing tag: {str(e)}")
            return False
            
    def get_document_tags(self, doc_id: int):
        """Get all tags for a document"""
        try:
            conn = self.conn_pool.get_connection()
            cursor = conn.cursor()
            cursor.execute(
                'SELECT tag_name FROM document_tags WHERE document_id = ? ORDER BY tag_name',
                (doc_id,)
            )
            return [row[0] for row in cursor.fetchall()]
        except Exception as e:
            logger.error(f"Error getting document tags: {str(e)}")
            return []
            
    def get_all_tags(self):
        """Get all unique tags in the system"""
        try:
            conn = self.conn_pool.get_connection()
            cursor = conn.cursor()
            cursor.execute(
                'SELECT DISTINCT tag_name FROM document_tags ORDER BY tag_name'
            )
            return [row[0] for row in cursor.fetchall()]
        except Exception as e:
            logger.error(f"Error getting all tags: {str(e)}")
            return []

    def backup_database(self, output_path: Union[str, Path]) -> str:
        """Create a transactionally consistent SQLite backup while the app runs."""
        output_path = Path(output_path)
        if output_path.resolve() == Path(self.db_path).resolve():
            raise ValueError("Backup path must differ from the active database")
        output_path.parent.mkdir(parents=True, exist_ok=True)
        # backup() drives its own transaction; only the handle needs closing,
        # otherwise the backup file stays locked on Windows.
        with closing(sqlite3.connect(output_path)) as destination:
            self.conn_pool.get_connection().backup(destination)
        return str(output_path)

    def close(self):
        """Close all database connections"""
        self.conn_pool.close_all()
